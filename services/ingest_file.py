"""文件接入：CSV / Excel / 邮件附件（readme 16.4）。

**核心断言：同一份数据走不同路径，落到 bronze 后必须等价。**
这条比任何单条路径的正确性都重要——它验证接入层没有引入偏差。

Excel 是脏数据的天然来源（合并单元格、多 sheet、标题行不在第一行、
日期被识别成数字、数字被存成文本），这些不用注入，真实文件自带。
"""
import csv
import hashlib
import os
import pathlib

MAX_BYTES = int(os.environ.get("FILE_MAX_BYTES", str(50 * 1024 * 1024)))
ALLOWED_EXT = {".csv", ".xlsx", ".xls", ".tsv"}


class FileRejected(ValueError):
    pass


def _guard(path: pathlib.Path):
    if not path.exists():
        raise FileRejected(f"文件不存在: {path}")
    if path.suffix.lower() not in ALLOWED_EXT:
        raise FileRejected(f"不支持的类型 {path.suffix}（仅 {sorted(ALLOWED_EXT)}）")
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise FileRejected(f"文件 {size:,} 字节超过上限 {MAX_BYTES:,}")
    return size


def _sniff_dialect(sample: str):
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        return csv.excel


def read_csv(path) -> dict:
    p = pathlib.Path(path)
    size = _guard(p)
    raw = p.read_bytes()
    # 编码嗅探：中文 CSV 常见 GBK
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileRejected("无法识别编码")
    dialect = _sniff_dialect(text[:4096])
    rows = list(csv.reader(text.splitlines(), dialect))
    if not rows:
        raise FileRejected("空文件")
    return {"source": "csv", "path": str(p), "bytes": size, "encoding": enc,
            "columns": rows[0], "rows": rows[1:]}


def read_excel(path, sheet=None) -> dict:
    """读 Excel。

    **不执行宏**：openpyxl 是纯解析库，不走 Excel 应用。
    标题行不在第一行是常见情况——这里取第一个非空行作为表头，
    但把跳过的行数报出来，让人能发现异常。
    """
    import openpyxl
    p = pathlib.Path(path)
    size = _guard(p)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    all_rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()

    skipped = 0
    while all_rows and not any(str(c).strip() for c in all_rows[0]):
        all_rows.pop(0)
        skipped += 1
    if not all_rows:
        raise FileRejected("空表")
    return {"source": "excel", "path": str(p), "bytes": size,
            "sheet": ws.title, "sheets": wb.sheetnames if hasattr(wb, "sheetnames") else [],
            "skipped_leading_blank_rows": skipped,
            "columns": [str(c) for c in all_rows[0]],
            "rows": [[c for c in r] for r in all_rows[1:]]}


def read_any(path, **kw) -> dict:
    ext = pathlib.Path(path).suffix.lower()
    return read_excel(path, **kw) if ext in (".xlsx", ".xls") else read_csv(path)


# ---------------------------------------------------------------- 等价性
def canonical_fingerprint(payload: dict) -> str:
    """内容指纹：用于验证「不同路径落到 bronze 等价」。

    归一化处理跨路径的固有差异：
    - Excel 把数字读成 int/float，CSV 全是字符串
    - 末尾空列、首尾空白
    这些是**格式差异不是数据差异**，归一化后必须一致。
    """
    def norm(v):
        s = str(v).strip()
        if s.endswith(".0") and s[:-2].lstrip("-").isdigit():
            s = s[:-2]                       # 1.0 -> 1
        return s

    cols = [norm(c) for c in payload["columns"]]
    while cols and not cols[-1]:
        cols.pop()
    lines = ["\x1f".join(cols)]
    for r in payload["rows"]:
        cells = [norm(c) for c in r][:len(cols)]
        cells += [""] * (len(cols) - len(cells))
        lines.append("\x1f".join(cells))
    return hashlib.sha256("\x1e".join(lines).encode()).hexdigest()


def profile_payload(payload: dict) -> dict:
    """对文件内容做与数据库表同口径的画像（复用 5.3 的指标）。"""
    cols = payload["columns"]
    n = len(payload["rows"])
    out = []
    for i, c in enumerate(cols):
        vals = [str(r[i]).strip() if i < len(r) else "" for r in payload["rows"]]
        nonnull = sum(1 for v in vals if v != "")
        distinct = len(set(v for v in vals if v != ""))
        out.append({"column": c,
                    "null_rate": round(1 - nonnull / n, 4) if n else 0,
                    "distinct": distinct,
                    "looks_unique": distinct == n and n > 1})
    return {"rows": n, "columns": out}


def read_table(source_id: str, table: str) -> dict:
    """第三条路径：数据库直连。与文件路径产出同一种 payload，便于比指纹。"""
    import connector
    import re
    if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", table or ""):
        raise FileRejected(f"表名不合法: {table!r}")
    r = connector.query(source_id, f'SELECT * FROM "{table}"', purpose="ingest")
    return {"source": "db", "table": table,
            "columns": list(r["columns"]),
            "rows": [list(row) for row in r["rows"]]}
